import os
import subprocess
import shutil
from typing import Dict, Any, Optional, List
from app.services.processor import DocumentProcessor, ProcessingError
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

class WordToPdfProcessor(DocumentProcessor):
    def __init__(self):
        super().__init__("word-to-pdf")

    def _convert_via_libreoffice(self, docx_path: str, output_path: str) -> bool:
        soffice_bin = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice_bin:
            return False
        
        out_dir = os.path.dirname(output_path)
        cmd = [soffice_bin, "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path]
        try:
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
            if res.returncode == 0:
                base_name = os.path.splitext(os.path.basename(docx_path))[0]
                generated_pdf = os.path.join(out_dir, f"{base_name}.pdf")
                if os.path.exists(generated_pdf):
                    if generated_pdf != output_path:
                        shutil.move(generated_pdf, output_path)
                    return True
        except Exception:
            pass
        return False

    def _convert_via_reportlab(self, docx_path: str, output_path: str):
        doc = docx.Document(docx_path)
        pdf_doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
        )
        
        styles = getSampleStyleSheet()
        normal_style = styles['Normal']
        normal_style.fontSize = 11
        normal_style.leading = 14

        heading1_style = ParagraphStyle(
            'Heading1',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            spaceAfter=10,
            textColor=colors.HexColor("#1e293b")
        )
        
        story = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                story.append(Spacer(1, 8))
                continue
            
            if p.style.name.startswith('Heading 1'):
                story.append(Paragraph(text, heading1_style))
                story.append(Spacer(1, 6))
            elif p.style.name.startswith('Heading'):
                h_style = ParagraphStyle('SubHeading', parent=heading1_style, fontSize=14, leading=18)
                story.append(Paragraph(text, h_style))
                story.append(Spacer(1, 4))
            else:
                # Escape xml tags in paragraph text
                safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe_text, normal_style))
                story.append(Spacer(1, 6))

        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    cell_text = cell.text.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    row_data.append(Paragraph(cell_text, normal_style))
                table_data.append(row_data)
            
            if table_data:
                t = Table(table_data)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#f1f5f9")),
                    ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor("#0f172a")),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
                ]))
                story.append(t)
                story.append(Spacer(1, 10))

        if not story:
            story.append(Paragraph("Document had no extractable paragraph text.", normal_style))

        pdf_doc.build(story)

    def process(self, input_paths: List[str], options: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        if not input_paths:
            raise ProcessingError("No input DOCX file provided", "INVALID_INPUT")

        docx_path = input_paths[0]
        if not os.path.exists(docx_path):
            raise ProcessingError("Input DOCX file does not exist", "FILE_NOT_FOUND")

        output_path = self.prepare_output_file("pdf")

        # Try LibreOffice first if available, else fallback to native reportlab engine
        success = self._convert_via_libreoffice(docx_path, output_path)
        if not success:
            try:
                self._convert_via_reportlab(docx_path, output_path)
            except Exception as e:
                raise ProcessingError(f"Word to PDF conversion failed: {str(e)}", "CONVERSION_ERROR")

        if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
            raise ProcessingError("Failed to generate output PDF document.", "EMPTY_OUTPUT")

        return {
            "output_path": output_path,
            "output_filename": "converted.pdf",
            "file_size": os.path.getsize(output_path),
        }
