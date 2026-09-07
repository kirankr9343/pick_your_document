import os
from PIL import Image, ImageDraw, ImageFont
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

def create_synthetic_test_data():
    out_dir = os.path.join(os.path.dirname(__file__), "..", "test-data")
    os.makedirs(out_dir, exist_ok=True)

    # 1. sample.jpg & sample.png
    img = Image.new("RGB", (600, 300), color=(240, 243, 246))
    draw = ImageDraw.Draw(img)
    draw.rectangle([20, 20, 580, 280], outline=(40, 100, 220), width=3)
    draw.text((40, 50), "Pick Your Document OCR Test Image", fill=(20, 30, 50))
    draw.text((40, 100), "Sample Extracted Text Line 1", fill=(30, 40, 60))
    draw.text((40, 140), "Sample Extracted Text Line 2 - 1234567890", fill=(30, 40, 60))
    img.save(os.path.join(out_dir, "sample.jpg"), "JPEG")
    img.save(os.path.join(out_dir, "sample.png"), "PNG")

    # 2. sample.docx
    doc = docx.Document()
    doc.add_heading("Pick Your Document - Synthetic Test File", level=1)
    doc.add_paragraph("This is a sample Word document generated for unit testing document conversion pipelines.")
    p = doc.add_paragraph("Features include paragraphs, headings, and structured text tables.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"
    doc.save(os.path.join(out_dir, "sample.docx"))

    # 3. sample.pdf
    styles = getSampleStyleSheet()
    pdf1 = SimpleDocTemplate(os.path.join(out_dir, "sample.pdf"), pagesize=letter)
    story1 = [
        Paragraph("Pick Your Document - Sample Single Page PDF", styles['Heading1']),
        Spacer(1, 12),
        Paragraph("This synthetic PDF document contains sample text for testing PDF to Word, PDF to Text, Compress PDF, and AI Summarization.", styles['Normal']),
        Spacer(1, 12),
        Paragraph("Key Point: All conversion functions must process clean data accurately without data corruption.", styles['Normal'])
    ]
    pdf1.build(story1)

    # 4. sample-multi-page.pdf
    from reportlab.platypus import PageBreak
    pdf_multi = SimpleDocTemplate(os.path.join(out_dir, "sample-multi-page.pdf"), pagesize=letter)
    story_multi = [
        Paragraph("Multi-Page PDF Document - Page 1", styles['Heading1']),
        Spacer(1, 12),
        Paragraph("First page content for testing PDF Split, PDF Merge, and PDF to JPG multi-page extraction.", styles['Normal']),
        PageBreak(),
        Paragraph("Multi-Page PDF Document - Page 2", styles['Heading1']),
        Spacer(1, 12),
        Paragraph("Second page content for testing range extraction (pages 1-2).", styles['Normal'])
    ]
    pdf_multi.build(story_multi)

    print("Synthetic test files created successfully in test-data/.")

if __name__ == "__main__":
    create_synthetic_test_data()
